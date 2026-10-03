import csv
import io
from pathlib import Path

from docx import Document as WordDocument
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".csv", ".txt", ".json"}
MAX_FILE_BYTES = 25 * 1024 * 1024


def extract_document(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Return ordered text sections with one-based page/slide references when available."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported yet.")
        return [
            (index + 1, page.extract_text(extraction_mode="layout") or "")
            for index, page in enumerate(reader.pages)
        ]
    if extension == ".docx":
        document = WordDocument(io.BytesIO(data))
        rows = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            rows.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return [(None, "\n".join(rows))]
    if extension == ".pptx":
        presentation = Presentation(io.BytesIO(data))
        slides = []
        for index, slide in enumerate(presentation.slides):
            parts = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
            slides.append((index + 1, "\n".join(parts)))
        return slides
    if extension == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="replace")
        rows = [" | ".join(cell.strip() for cell in row) for row in csv.reader(io.StringIO(text))]
        return [(None, "\n".join(rows))]
    if extension == ".txt":
        text = data.decode("utf-8", errors="replace")
        return [(None, text)]
    if extension == ".json":
        text = data.decode("utf-8", errors="replace")
        return [(None, text)]
    raise ValueError("Supported formats are PDF, DOCX, PPTX, CSV, TXT, and JSON.")

